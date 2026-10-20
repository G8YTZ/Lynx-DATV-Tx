#!/usr/bin/env python3
"""Builds the Lynx-Tx Word guides (docs/*.docx) from docs/src/*.md.

Same look as the Lynx DATV Receiver's guides: docs/src/template.docx is the
Receiver's install guide with its text removed, so the styles, page size and
bullets are the same. The sources use a small subset of Markdown:

  title: / subtitle: / tagline: / version:   front matter, first lines
  # ## ###                                   headings
  - item                                     bullets (one level)
  ```  ...  ```                              command block
  > text                                     note (amber box)
  | a | b |                                  table, first row the header
  ![alt](img/file.png){5.5}                  picture, width in inches
  **bold**  *italic*  `code`  [text](url)    inline

Run from anywhere:  python3 tools/build_docs.py
"""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs")
SRC = os.path.join(DOCS, "src")
TEMPLATE = os.path.join(SRC, "template.docx")

CODE_FONT = "Consolas"
HEADER_FILL = "2C3E50"


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def spacing(p, before=None, after=None):
    ppr = p._p.get_or_add_pPr()
    s = el("w:spacing")
    if before is not None:
        s.set(qn("w:before"), str(before))
    if after is not None:
        s.set(qn("w:after"), str(after))
    ppr.append(s)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*\s][^*]*\*(?!\*))")


def add_hyperlink(p, text, url, size=None, bold=False):
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = el("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = el("w:r")
    rpr = el("w:rPr")
    rpr.append(el("w:color", **{"w:val": "2E74B5"}))
    rpr.append(el("w:u", **{"w:val": "single"}))
    if bold:
        rpr.append(el("w:b"))
    if size:
        rpr.append(el("w:sz", **{"w:val": size}))
    r.append(rpr)
    t = el("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def add_inline(p, text, size=None, bold=False, italic=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, code = bold, italic, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
            # a link or code inside bold
            if INLINE.search(part):
                add_inline(p, part, size, True, italic, color)
                continue
        elif part.startswith("`") and part.endswith("`"):
            part, code = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            add_hyperlink(p, m.group(1), m.group(2), size, bold)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        run = p.add_run(part)
        run.bold = b or None
        run.italic = i or None
        if code:
            run.font.name = CODE_FONT
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), CODE_FONT)
        if size:
            run.font.size = Pt(size / 2)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def para(doc, text, after=150):
    p = doc.add_paragraph()
    spacing(p, after=after)
    add_inline(p, text)
    return p


def bullet(doc, text):
    p = styled(doc, "ListParagraph")
    ppr = p._p.get_or_add_pPr()
    num = el("w:numPr")
    num.append(el("w:ilvl", **{"w:val": 0}))
    num.append(el("w:numId", **{"w:val": 2}))
    ppr.append(num)
    spacing(p, after=80)
    add_inline(p, text)


def code(doc, lines):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    ppr.append(el("w:shd", **{"w:fill": "F0F0F0", "w:val": "clear"}))
    spacing(p, before=100, after=150)
    ppr.append(el("w:ind", **{"w:left": 360}))
    for n, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = CODE_FONT
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), CODE_FONT)
        run.font.size = Pt(10)
        if n < len(lines) - 1:
            run.add_break()


def note(doc, text):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    bdr = el("w:pBdr")
    bdr.append(el("w:left", **{"w:val": "single", "w:color": "E8A33D", "w:sz": 12, "w:space": 8}))
    ppr.append(bdr)
    ppr.append(el("w:shd", **{"w:fill": "FFF4E5", "w:val": "clear"}))
    spacing(p, before=100, after=200)
    ppr.append(el("w:ind", **{"w:left": 200}))
    add_inline(p, text, size=20, italic=True)


def table(doc, rows):
    total = 9360  # text width, twentieths of a point (6.5 inches)
    ncol = len(rows[0])
    # Each column at least as wide as its longest word (so nothing breaks
    # mid-word), the rest shared out by how much text each column holds.
    def plain(x):
        return re.sub(r"[*`]|\[([^\]]+)\]\([^)]+\)", r"\1", x)
    cols = [[plain(r[c]) if c < len(r) else "" for r in rows] for c in range(ncol)]
    minw = [max(760, 110 * max((len(w) for x in col for w in x.split()), default=3) + 200)
            for col in cols]
    if sum(minw) > total:
        minw = [int(m * total / sum(minw)) for m in minw]
    weight = [min(max(sum(len(x) for x in col) / len(col), 4), 60) for col in cols]
    spare = total - sum(minw)
    widths = [m + int(spare * w / sum(weight)) for m, w in zip(minw, weight)]
    t = doc.add_table(rows=0, cols=ncol)
    tbl = t._tbl
    tblpr = tbl.tblPr
    tblpr.append(el("w:tblW", **{"w:type": "dxa", "w:w": total}))
    borders = el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(el(f"w:{side}", **{"w:val": "single", "w:color": "auto", "w:sz": 4}))
    tblpr.append(borders)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            c = cells[ci]
            tcpr = c._tc.get_or_add_tcPr()
            tcpr.append(el("w:tcW", **{"w:type": "dxa", "w:w": widths[ci]}))
            if ri == 0:
                tcpr.append(el("w:shd", **{"w:fill": HEADER_FILL, "w:val": "clear"}))
            p = c.paragraphs[0]
            text = row[ci] if ci < len(row) else ""
            if ri == 0:
                add_inline(p, text, size=20, bold=True, color="FFFFFF")
            else:
                add_inline(p, text, size=20)
    # Header row repeats on each page.
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(el("w:tblHeader"))
    gap = doc.add_paragraph()
    spacing(gap, after=60)


def image(doc, path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, before=100, after=150)
    p.add_run().add_picture(path, width=Inches(width))


def styled(doc, style_id, text=""):
    """A paragraph in one of the template's styles, by its id: the
    Receiver's template names them "Heading 1" and so on, which python-docx's
    name lookup does not find."""
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().insert(0, el("w:pStyle", **{"w:val": style_id}))
    if text:
        p.add_run(text)
    return p


def heading(doc, text, level):
    styled(doc, f"Heading{level}", text)


def build(src):
    with open(src, encoding="utf-8") as f:
        lines = f.read().split("\n")
    meta = {}
    while lines and re.match(r"^(title|subtitle|tagline|version):", lines[0]):
        k, v = lines.pop(0).split(":", 1)
        meta[k] = v.strip()
    doc = Document(TEMPLATE)
    doc.core_properties.title = f"{meta.get('title', 'Lynx')} {meta.get('subtitle', '')}".strip()
    doc.core_properties.author = "Justin Cockett, G8YTZ"
    styled(doc, "Title", meta.get("title", "Lynx"))
    if "subtitle" in meta:
        heading(doc, meta["subtitle"], 1)
    if "tagline" in meta:
        para(doc, meta["tagline"], after=100)
    if "version" in meta:
        p = doc.add_paragraph()
        spacing(p, after=150)
        p.add_run(meta["version"]).bold = True

    i, buf = 0, []

    def flush():
        if buf:
            para(doc, " ".join(s.strip() for s in buf))
            buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            j, block = i + 1, []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j].rstrip())
                j += 1
            code(doc, block)
            i = j + 1
            continue
        m = re.match(r"^(#{1,3}) (.+)$", s)
        if m:
            flush()
            heading(doc, m.group(2), len(m.group(1)))
        elif s.startswith("- "):
            flush()
            text = s[2:]
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") and lines[i + 1].strip() \
                    and not lines[i + 1].strip().startswith(("- ", "```", "|")):
                i += 1
                text += " " + lines[i].strip()
            bullet(doc, text)
        elif s.startswith("> "):
            flush()
            text = s[2:]
            while i + 1 < len(lines) and lines[i + 1].strip().startswith("> "):
                i += 1
                text += " " + lines[i].strip()[2:]
            note(doc, text)
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue
        elif s.startswith("!["):
            flush()
            m = re.match(r"!\[[^\]]*\]\(([^)]+)\)(?:\{([\d.]+)\})?", s)
            image(doc, os.path.join(DOCS, m.group(1)), float(m.group(2) or 6))
        elif not s:
            flush()
        else:
            buf.append(s)
        i += 1
    flush()
    out = os.path.join(DOCS, os.path.splitext(os.path.basename(src))[0] + ".docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    srcs = sys.argv[1:] or sorted(
        os.path.join(SRC, f) for f in os.listdir(SRC) if f.endswith(".md"))
    for s in srcs:
        print("built", os.path.relpath(build(s)))
